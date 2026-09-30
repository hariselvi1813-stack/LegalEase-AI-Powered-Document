"""
LegalEase Backend - FastAPI Application
Provides document generation and export endpoints with strict error boundaries.
"""

import os
import io
from typing import List, Optional
from fastapi import FastAPI, HTTPException, status, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from dotenv import load_dotenv

# Load local environment if available
load_dotenv()

from backend.services import generate_with_gemini, build_demo_document

app = FastAPI(
    title="LegalEase API",
    description="Beginner-friendly AI legal document generator on Free Tiers",
    version="1.0.0"
)

# Enable CORS for Streamlit and web frontends
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# Request Models
class DocumentRequest(BaseModel):
    doc_type: str = Field(..., description="Document type, e.g., NDA, Employment Contract, Lease Agreement, Freelance Contract, Custom")
    parties: List[str] = Field(default_factory=list, description="List of participating party names")
    terms: List[str] = Field(default_factory=list, description="Key clauses or terms to include")
    effective_date: str = Field(..., description="Effective date string, e.g., '2026-10-01'")


class ExportRequest(BaseModel):
    title: str = Field(default="Legal_Document", description="Title of the document")
    content: str = Field(..., description="Full text content of the legal document")


# Endpoints
@app.get("/health")
def health_check():
    """Health check endpoint to verify backend availability."""
    return {"status": "ok"}


@app.post("/generate")
def generate_document(req: DocumentRequest):
    """
    Generates a legal document via Gemini or Demo Mode.
    Guaranteed never to crash or expose internal secrets.
    """
    # Validation
    if not req.doc_type or not req.doc_type.strip():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Document type is required.")
    
    if len(req.doc_type) > 100:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Document type is too long.")

    if len(req.parties) > 20:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Too many parties specified.")

    # Generate document
    try:
        content, mode = generate_with_gemini(
            doc_type=req.doc_type.strip(),
            parties=req.parties,
            terms=req.terms,
            effective_date=req.effective_date.strip() if req.effective_date else "Effective Immediately"
        )
        return {"content": content, "mode": mode}
    except Exception:
        # Fallback to demo mode on unexpected error
        demo_content = build_demo_document(
            doc_type=req.doc_type,
            parties=req.parties,
            terms=req.terms,
            effective_date=req.effective_date
        )
        return {"content": demo_content, "mode": "demo"}


@app.post("/generate-docx")
def generate_docx(req: ExportRequest):
    """
    Builds and downloads a DOCX file from the document content.
    """
    if not req.content or not req.content.strip():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Document content cannot be empty.")

    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor

        doc = Document()
        
        # Configure margins
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # Title
        title_p = doc.add_paragraph()
        run = title_p.add_run(req.title.replace("_", " ").upper())
        run.bold = True
        run.font.name = "Georgia"
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(11, 31, 58) # Deep navy

        # Body paragraphs
        lines = req.content.split("\n")
        for line in lines:
            line_str = line.strip()
            if not line_str:
                doc.add_paragraph("")
                continue
            
            p = doc.add_paragraph()
            p_run = p.add_run(line_str)
            p_run.font.name = "Georgia"
            p_run.font.size = Pt(11)
            p_run.font.color.rgb = RGBColor(30, 41, 59)

        # Save to buffer
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        filename = f"{req.title.strip().replace(' ', '_')}.docx"
        return Response(
            content=file_stream.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Unable to generate the document. Please try again."
        )


@app.post("/generate-pdf")
def generate_pdf(req: ExportRequest):
    """
    Builds and downloads a PDF file from the document content using fpdf.
    """
    if not req.content or not req.content.strip():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Document content cannot be empty.")

    try:
        from fpdf import FPDF

        class LegalPDF(FPDF):
            def header(self):
                self.set_font("Helvetica", "B", 10)
                self.set_text_color(120, 120, 120)
                self.cell(0, 10, "LegalEase - Confidential Legal Document", 0, 1, "R")
                self.ln(2)

            def footer(self):
                self.set_y(-15)
                self.set_font("Helvetica", "I", 8)
                self.set_text_color(150, 150, 150)
                self.cell(0, 10, f"Page {self.page_no()}", 0, 0, "C")

        pdf = LegalPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_font("Times", size=11)
        pdf.set_text_color(20, 20, 20)

        # Write lines
        lines = req.content.split("\n")
        for line in lines:
            safe_text = line.encode("latin-1", "replace").decode("latin-1")
            pdf.multi_cell(0, 6, safe_text)

        pdf_bytes = pdf.output(dest="S").encode("latin-1")
        filename = f"{req.title.strip().replace(' ', '_')}.pdf"
        
        return Response(
            content=pdf_bytes,
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Unable to generate the document. Please try again."
        )
